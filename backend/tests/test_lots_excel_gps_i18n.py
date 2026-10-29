"""
Cropbook - Lots Excel export: separate Location / GPS columns, GPS hyperlink, translated type/product (it/en/es)
"""
import io
import os
import time
import pytest
import requests
import openpyxl

BASE = os.environ.get("REACT_APP_BACKEND_URL", "https://app-review-56.preview.emergentagent.com").rstrip("/")
API = f"{BASE}/api"
TS = int(time.time())
EXPECT = {
    "it": {"loc": "Luogo", "gps": "Coordinate GPS", "map": "Apri mappa", "type": "Frutta", "prod": "Mele", "olive": "Olivo"},
    "en": {"loc": "Location", "gps": "GPS Coordinates", "map": "Open map", "type": "Fruit", "prod": "Apples", "olive": "Olive"},
    "es": {"loc": "Lugar", "gps": "Coordenadas GPS", "map": "Abrir mapa", "type": "Fruta", "prod": "Manzanas", "olive": "Olivo"},
}


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{API}/auth/login", json={"username": "admin", "password": "9ded04059b99"}, timeout=20)
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def lots(token):
    h = {"Authorization": f"Bearer {token}"}
    created = {}
    specs = {
        "maps": {"gps_coordinates": "https://maps.google.com/?q=45.4765,10.9257", "product_type": "frutta", "product_category": "Mele"},
        "short": {"gps_coordinates": "https://maps.app.goo.gl/abcDEF123", "product_type": "olivo", "product_category": "altri"},
        "at": {"gps_coordinates": "https://www.google.com/maps/@41.9028,12.4964,15z", "product_type": "frutta", "product_category": "Arance"},
    }
    for k, s in specs.items():
        body = {"company_name": f"TEST_XLS_{TS}", "location": f"TEST_Loc_{k}_{TS}", "field_size": 1.5,
                "field_lot": f"TEST_FL_{k}", **s}
        r = requests.post(f"{API}/lots", json=body, headers=h, timeout=20)
        assert r.status_code in (200, 201), f"{k}: {r.status_code} {r.text}"
        d = r.json()
        created[k] = (d.get("data") or {}).get("id") or d.get("id")
        assert created[k], d
    yield created
    for lid in created.values():
        requests.delete(f"{API}/lots/{lid}", headers=h, timeout=20)
    # cleanup leftovers from earlier runs + test companies
    r = requests.get(f"{API}/lots", headers=h, timeout=20)
    items = r.json().get("data", r.json()) if r.ok else []
    for l in items if isinstance(items, list) else []:
        if str(l.get("location", "")).startswith("TEST_Loc_"):
            requests.delete(f"{API}/lots/{l['id']}", headers=h, timeout=20)
    r = requests.get(f"{API}/companies", headers=h, timeout=20)
    items = r.json().get("data", r.json()) if r.ok else []
    for c in items if isinstance(items, list) else []:
        if str(c.get("name", "")).startswith("TEST_XLS_"):
            requests.delete(f"{API}/companies/{c['id']}", headers=h, timeout=20)


def get_sheet(token, lang):
    r = requests.get(f"{API}/lots/export/excel", params={"lang": lang},
                     headers={"Authorization": f"Bearer {token}"}, timeout=60)
    assert r.status_code == 200, r.text[:300]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    return wb.active


def row_by_location(ws, headers, loc):
    li = headers.index
    for row in ws.iter_rows(min_row=2):
        if row[li(EXP_LOC[0])].value == loc:
            return row
    return None


EXP_LOC = [None]


@pytest.mark.parametrize("lang", ["it", "en", "es"])
def test_excel_columns_and_gps(token, lots, lang):
    e = EXPECT[lang]
    ws = get_sheet(token, lang)
    headers = [c.value for c in ws[1]]
    print(lang, headers)
    assert e["loc"] in headers and e["gps"] in headers
    assert headers.index(e["gps"]) == headers.index(e["loc"]) + 1
    li, gi = headers.index(e["loc"]), headers.index(e["gps"])
    rows = {r[li].value: r for r in ws.iter_rows(min_row=2)}

    r = rows[f"TEST_Loc_maps_{TS}"]
    assert r[gi].value == "45.476500, 10.925700", r[gi].value
    assert r[gi].hyperlink is not None and r[gi].hyperlink.target == "https://maps.google.com/?q=45.4765,10.9257"
    assert r[1].value == e["type"] and r[2].value == e["prod"], (r[1].value, r[2].value)

    r = rows[f"TEST_Loc_short_{TS}"]
    assert r[gi].value == e["map"], r[gi].value
    assert r[gi].hyperlink.target == "https://maps.app.goo.gl/abcDEF123"
    assert r[1].value == e["olive"]
    print(lang, "altri ->", r[2].value)

    r = rows[f"TEST_Loc_at_{TS}"]
    assert r[gi].value == "41.902800, 12.496400", r[gi].value
    assert r[gi].hyperlink.target.startswith("https://www.google.com/maps/@41.9028")
